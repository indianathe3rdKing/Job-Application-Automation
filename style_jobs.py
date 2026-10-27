

from pathlib import Path
import sys
import subprocess
from openpyxl import load_workbook
from openpyxl.styles import (
    Font, Alignment, PatternFill, Border, Side
)
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

INPUT_FILE = Path("jobs.xlsx")
OUTPUT_FILE = Path("jobs_styled.xlsx")

if not INPUT_FILE.exists():
    raise FileNotFoundError(
        f"Could not find {INPUT_FILE}. "
        "Check the filename and folder."
    )

workbook = load_workbook(INPUT_FILE)

header_fill = PatternFill("solid", fgColor="172554")
header_font = Font(
    name="Aptos",
    bold=True,
    color="FFFFFF",
    size=11,
)
header_alignment = Alignment(
    horizontal="left",
    vertical="center",
    wrap_text=True,
)

thin_border = Border(
    bottom=Side(style="thin", color="E2E8F0")
)

for sheet_index, sheet in enumerate(workbook.worksheets, start=1):
    if sheet.max_row < 1 or sheet.max_column < 1:
        continue

    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    sheet.row_dimensions[1].height = 32

    # Style headers
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment

    # Set column widths
    for column_cells in sheet.columns:
        column_letter = get_column_letter(
            column_cells[0].column
        )
        header = str(
            column_cells[0].value or ""
        ).lower()

        if "title" in header:
            width = 36
        elif "company" in header:
            width = 26
        elif "location" in header:
            width = 28
        elif "url" in header:
            width = 24
        else:
            max_length = max(
                (
                    len(str(cell.value))
                    for cell in column_cells
                    if cell.value is not None
                ),
                default=10,
            )
            width = max_length + 3

        sheet.column_dimensions[column_letter].width = min(
            max(width, 14), 45
        )

    # Style rows and make job URLs clickable
    for row in range(2, sheet.max_row + 1):
        sheet.row_dimensions[row].height = 26

        for cell in sheet[row]:
            cell.alignment = Alignment(
                vertical="center",
                wrap_text=False,
            )
            cell.border = thin_border

        # Find the URL column by its header
        url_column = next(
            (
                cell.column
                for cell in sheet[1]
                if str(cell.value or "").lower()
                in ("job_url", "job url")
            ),
            None,
        )

        if url_column is not None:
            url_cell = sheet.cell(
                row=row,
                column=url_column,
            )
            url = url_cell.value

            if isinstance(url, str) and url.startswith(
                ("https://", "http://")
            ):
                url_cell.hyperlink = url
                url_cell.style = "Hyperlink"
                url_cell.value = "Open job listing"

    # Add one formatted Excel table per worksheet
    if sheet.max_row >= 2:
        table = Table(
            displayName=f"JobListings{sheet_index}",
            ref=sheet.dimensions,
        )

        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showRowStripes=True,
            showColumnStripes=False,
            showFirstColumn=False,
            showLastColumn=False,
        )

        sheet.add_table(table)

workbook.save(OUTPUT_FILE)

print(f"Created {OUTPUT_FILE.resolve()} successfully!")

subprocess.run(
    [sys.executable, "process_job.py"],
    check=True
)
