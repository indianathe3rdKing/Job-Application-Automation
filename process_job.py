from copy import copy
from pathlib import Path
import re

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries
from openpyxl.worksheet.table import TableColumn

INPUT_FILE = Path("jobs_styled.xlsx")
OUTPUT_FILE = Path("jobs_ranked.xlsx")


ROLE_KEYWORDS = {
    "Android / Kotlin": [
        "android developer", "kotlin developer", "android engineer",
        "mobile developer", "mobile engineer"
    ],
    "Frontend / JavaScript": [
        "frontend developer", "front-end developer", "react developer",
        "javascript developer", "web developer", "react native developer"
    ],
    "Software Development": [
        "junior software developer", "software developer",
        "graduate developer", "python developer", "java developer",
        "application developer", "software engineer"
    ],
    "Internships / Graduate": [
        "software intern", "developer intern", "development internship",
        "graduate programme", "graduate program", "graduate trainee"
    ],
    "Cloud / AWS": [
        "cloud support", "junior cloud", "aws support",
        "cloud operations", "cloud engineer", "cloud administrator"
    ],
    "IT / Application Support": [
        "technical support", "application support", "it support",
        "software support", "service desk", "help desk",
        "desktop support", "it technician", "systems support"
    ],
    "QA / Testing": [
        "qa tester", "qa analyst", "software tester",
        "manual tester", "junior test analyst", "quality assurance"
    ],
    "Business / Product Support": [
        "product support", "implementation specialist",
        "implementation support", "customer success",
        "technical account", "client onboarding",
        "business analyst", "junior business analyst"
    ],
    "Customer Service": [
        "customer service", "customer support", "customer care",
        "call centre", "call center", "customer experience",
        "client services", "customer operations"
    ],
    "Sales / Technical Sales": [
        "technical sales", "sales support", "inside sales",
        "sales development representative", "sales representative",
        "account representative", "b2b sales"
    ],
    "Administration / Operations": [
        "administrative assistant", "office administrator",
        "operations administrator", "data capturer",
        "data entry", "reporting assistant", "crm administrator"
    ],
}

# Higher numbers mean greater relevance to your current career direction.
CATEGORY_WEIGHTS = {
    "Android / Kotlin": 65,
    "Frontend / JavaScript": 63,
    "Software Development": 62,
    "Internships / Graduate": 60,
    "Cloud / AWS": 55,
    "IT / Application Support": 55,
    "QA / Testing": 52,
    "Business / Product Support": 48,
    "Customer Service": 40,
    "Sales / Technical Sales": 38,
    "Administration / Operations": 30,
}

SENIOR_KEYWORDS = [
    "senior", "lead", "principal", "staff", "engineering lead",
    "software lead", "technical lead", "team lead", "engineering manager",
    "engineering team manager",
    "software engineering manager", "head of engineering",
    "head of software", "director of engineering",
    "director of software engineering", "director of technology",
    "vp of engineering", "vice president of engineering",
    "chief technology officer", "chief technical officer",
    "chief information officer", "cto", "director", "vp",
    "vice president", "chief",
]

SENIOR_DESCRIPTION_PATTERNS = [
    r"\bthis role is (?:a|an) (?:senior|lead|principal|staff)\b",
    r"\bthe position is (?:a|an) (?:senior|lead|principal|staff)\b",
    r"\b(?:you|the successful candidate|the role|this position) "
    r"(?:will |is expected to )?(?:manage|lead|oversee) "
    r"(?:the|a|our) (?:(?:software|engineering|development|technology) )?"
    r"team\b",
]

INDUSTRIAL_TITLE_KEYWORDS = [
    "boiler operator", "plant operator", "machine operator",
    "factory operator", "factory worker", "factory operative",
    "production operator", "production worker", "production operative",
    "production line operator", "production line worker",
    "production technician", "manufacturing operator",
    "manufacturing technician", "manufacturing engineer",
    "manufacturing associate", "industrial operator",
    "industrial technician", "industrial engineer",
    "industrial maintenance", "boiler technician", "plant technician",
    "machine technician", "machine setter", "process operator",
    "assembly line operator", "assembly operator", "assembly worker",
]

JUNIOR_KEYWORDS = [
    "junior", "entry level", "entry-level",
    "graduate", "intern", "trainee"
]

# Skills supported by your CV and projects.
CV_SKILLS = {
    "Kotlin": ["kotlin"],
    "Android": ["android", "jetpack compose"],
    "JavaScript": ["javascript", "ecmascript"],
    "React": ["react", "react.js"],
    "React Native": ["react native"],
    "Python": ["python"],
    "AWS": ["aws", "amazon web services"],
    "Lambda": ["lambda"],
    "DynamoDB": ["dynamodb"],
    "API Gateway": ["api gateway"],
    "REST APIs": ["rest api", "restful api"],
    "Git": ["git", "github"],
    "Appwrite": ["appwrite"],
    "Troubleshooting": [
        "troubleshooting", "technical diagnosis",
        "issue resolution", "problem resolution"
    ],
    "Customer Support": [
        "customer support", "technical support",
        "customer service", "client support"
    ],
    "Device Support": [
        "device troubleshooting", "device support",
        "hardware support", "payment device"
    ],
    "Onboarding": ["customer onboarding", "client onboarding"],
    "Sales": ["b2b sales", "sales targets", "sales representative"],
    "Team Leadership": [
        "team leader", "team leadership", "supervised a team"
    ],
    "IT Operations": [
        "software installation", "operating system updates",
        "data backups", "user account management"
    ],
    "Data Validation": ["data validation", "data quality"],
    "API Integration": ["api integration", "integrating apis"],
}

RESULT_COLUMNS = [
    "match_score",
    "match_category",
    "match_reason",
    "matched_skills",
    "application_status",
]

JUNIOR_BONUS = 20
SKILL_MATCH_BONUS = 3
MAX_SKILL_MATCH_BONUS = 15


def category_for_title(title):
    title = str(title).lower()
    for category, keywords in ROLE_KEYWORDS.items():
        if any(keyword in title for keyword in keywords):
            return category
    return None


def is_senior_role(title, description):
    title = str(title or "").lower()
    description = str(description or "").lower()
    if any(
        re.search(rf"\b{re.escape(keyword)}\b", title)
        for keyword in SENIOR_KEYWORDS
    ):
        return True
    return any(
        re.search(pattern, description)
        for pattern in SENIOR_DESCRIPTION_PATTERNS
    )


def is_industrial_role(title, description):
    title = str(title or "").lower()
    description = str(description or "").lower()
    if any(keyword in title for keyword in INDUSTRIAL_TITLE_KEYWORDS):
        return True

    has_industrial_context = any(
        keyword in description
        for keyword in (
            "manufacturing plant", "factory floor", "boiler room",
            "assembly line", "production line", "industrial machinery",
        )
    )
    has_industrial_job_title = any(
        keyword in title for keyword in ("operator", "operative", "worker")
    )
    return has_industrial_context and has_industrial_job_title


def analyze_description(description):
    description = str(description or "").lower()
    matched_skills = [
        skill
        for skill, keywords in CV_SKILLS.items()
        if any(keyword in description for keyword in keywords)
    ]
    return matched_skills


def score_job(title, matched_skills=()):
    category = category_for_title(title)
    if category is None:
        raise ValueError(f"Cannot score irrelevant job title: {title!r}")

    reasons = [f"Matches {category}"]
    score = CATEGORY_WEIGHTS[category]
    skill_bonus = min(
        len(matched_skills) * SKILL_MATCH_BONUS,
        MAX_SKILL_MATCH_BONUS,
    )
    if skill_bonus:
        score += skill_bonus
        reasons.append(f"CV skill overlap: {', '.join(matched_skills)}")

    normalized_title = str(title).lower()
    if any(
        re.search(rf"\b{re.escape(keyword)}\b", normalized_title)
        for keyword in JUNIOR_KEYWORDS
    ):
        score += JUNIOR_BONUS
        reasons.append("Entry-level indicator")

    return min(score, 100), category, "; ".join(reasons)


def normalized_header(value):
    return str(value or "").strip().lower().replace(" ", "_")


def table_for_header_row(worksheet):
    for table_name in worksheet.tables:
        table = worksheet.tables[table_name]
        if range_boundaries(table.ref)[1] == 1:
            return table
    return None


def process_worksheet(worksheet):
    legacy_gap_columns = [
        cell.column
        for cell in worksheet[1]
        if normalized_header(cell.value) == "missing_requirements"
    ]
    for column in reversed(legacy_gap_columns):
        worksheet.delete_cols(column)
    for table_name in worksheet.tables:
        table = worksheet.tables[table_name]
        if table.tableColumns:
            table.tableColumns = [
                column
                for column in table.tableColumns
                if normalized_header(column.name) != "missing_requirements"
            ]
            for index, column in enumerate(table.tableColumns, start=1):
                column.id = index

    header_columns = {
        normalized_header(cell.value): cell.column
        for cell in worksheet[1]
        if cell.value is not None
    }
    required_columns = {"title", "company", "location", "job_url"}
    missing_columns = required_columns - header_columns.keys()

    if missing_columns:
        return False

    result_columns = {}
    for column_name in RESULT_COLUMNS:
        column = header_columns.get(column_name)
        if column is None:
            column = worksheet.max_column + 1
            worksheet.cell(row=1, column=column, value=column_name)
            header_columns[column_name] = column
        result_columns[column_name] = column

    title_column = header_columns["title"]
    url_column = header_columns["job_url"]
    description_column = header_columns.get("description")
    max_column = worksheet.max_column
    records = []
    seen_urls = set()
    excluded_counts = {
        "duplicate": 0,
        "senior": 0,
        "industrial": 0,
        "irrelevant": 0,
    }

    for row_number in range(2, worksheet.max_row + 1):
        cells = [
            worksheet.cell(row=row_number, column=column)
            for column in range(1, max_column + 1)
        ]
        title = cells[title_column - 1].value
        url_cell = cells[url_column - 1]
        url = url_cell.hyperlink.target if url_cell.hyperlink else url_cell.value

        if title is None or url is None:
            continue

        title = str(title).strip()
        url_key = str(url).strip()
        if not title or not url_key:
            continue
        if url_key in seen_urls:
            excluded_counts["duplicate"] += 1
            continue
        seen_urls.add(url_key)

        description = (
            cells[description_column - 1].value
            if description_column is not None
            else ""
        )
        if is_senior_role(title, description):
            excluded_counts["senior"] += 1
            continue
        if is_industrial_role(title, description):
            excluded_counts["industrial"] += 1
            continue
        category = category_for_title(title)
        if category is None:
            excluded_counts["irrelevant"] += 1
            continue

        matched_skills = analyze_description(description)
        score, category, reason = score_job(title, matched_skills)
        values = [cell.value for cell in cells]
        values[title_column - 1] = title
        values[result_columns["match_score"] - 1] = score
        values[result_columns["match_category"] - 1] = category
        values[result_columns["match_reason"] - 1] = reason
        values[result_columns["matched_skills"] - 1] = ", ".join(
            matched_skills
        )
        values[result_columns["application_status"] - 1] = "Not applied"

        records.append(
            {
                "values": values,
                "styles": [copy(cell._style) for cell in cells],
                "hyperlinks": [
                    copy(cell.hyperlink) if cell.hyperlink else None
                    for cell in cells
                ],
                "score": score,
                "title": title,
                "company": cells[header_columns["company"] - 1].value,
                "category": category,
                "reason": reason,
            }
        )

    records.sort(key=lambda record: record["score"], reverse=True)

    header_style = copy(worksheet.cell(row=1, column=1)._style)
    row_style = copy(worksheet.cell(row=2, column=1)._style)
    header_alignment = copy(worksheet.cell(row=1, column=1).alignment)
    row_alignment = copy(worksheet.cell(row=2, column=1).alignment)
    row_border = copy(worksheet.cell(row=2, column=1).border)

    for row_number, record in enumerate(records, start=2):
        for column, value in enumerate(record["values"], start=1):
            cell = worksheet.cell(row=row_number, column=column)
            cell.value = value
            cell._style = copy(record["styles"][column - 1])
            cell.hyperlink = copy(record["hyperlinks"][column - 1])

    trailing_rows = worksheet.max_row - len(records) - 1
    if trailing_rows > 0:
        worksheet.delete_rows(len(records) + 2, trailing_rows)

    for column in result_columns.values():
        header_cell = worksheet.cell(row=1, column=column)
        header_cell._style = copy(header_style)
        header_cell.alignment = copy(header_alignment)

        for row_number in range(2, len(records) + 2):
            cell = worksheet.cell(row=row_number, column=column)
            cell._style = copy(row_style)
            cell.alignment = copy(row_alignment)
            cell.border = copy(row_border)

    table = table_for_header_row(worksheet)
    if table is not None:
        table.ref = f"A1:{worksheet.cell(row=1, column=worksheet.max_column).column_letter}{len(records) + 1}"
        if table.autoFilter is not None:
            table.autoFilter.ref = table.ref

        existing_table_columns = {
            column.name for column in table.tableColumns
        }
        for column_name in result_columns:
            if column_name not in existing_table_columns:
                table.tableColumns.append(
                    TableColumn(
                        id=len(table.tableColumns) + 1,
                        name=column_name,
                    )
                )

    for column_name, column in result_columns.items():
        width = {
            "match_score": 14,
            "match_category": 26,
            "match_reason": 38,
            "matched_skills": 32,
            "application_status": 18,
        }[column_name]
        worksheet.column_dimensions[
            worksheet.cell(row=1, column=column).column_letter
        ].width = width

    print(f"{worksheet.title}: ranked {len(records)} listings")
    print(
        f"Excluded: {excluded_counts['duplicate']} duplicates, "
        f"{excluded_counts['senior']} senior, "
        f"{excluded_counts['industrial']} industrial, "
        f"{excluded_counts['irrelevant']} irrelevant"
    )
    print("Top 10 matches")
    for record in records[:10]:
        print(
            f"{record['title']} | {record['company']} | "
            f"{record['category']} | {record['score']} | {record['reason']}"
        )
    return True


def main():
    if not INPUT_FILE.exists():
        raise FileNotFoundError(
            f"Could not find {INPUT_FILE}. Run scrape_jobs.py first."
        )

    workbook = load_workbook(INPUT_FILE)
    processed_sheets = [
        worksheet
        for worksheet in workbook.worksheets
        if process_worksheet(worksheet)
    ]

    if not processed_sheets:
        raise ValueError(
            f"No worksheet in {INPUT_FILE} contains the required "
            "title, company, location, and job_url columns."
        )

    workbook.save(OUTPUT_FILE)
    total_listings = sum(
        worksheet.max_row - 1 for worksheet in processed_sheets
    )
    print(f"Ranked spreadsheet saved to: {OUTPUT_FILE.resolve()}")
    print(f"Total listings ranked: {total_listings}")


if __name__ == "__main__":
    main()
